# 校验生成的 xlsx: zip 完整性 / XML well-formed / 关系一致 / openpyxl 可打开 / 公式 / 数字格式 / 冻结窗格
# 用法: python validate.py gen/out_all.xlsx gen/out_md.xlsx gen/out_edit.xlsx
import re
import sys
import zipfile
import json
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

def validate(path):
    print(f'== {path} ==')
    z = zipfile.ZipFile(path)
    assert z.testzip() is None, 'zip CRC 损坏'
    names = set(z.namelist())

    for n in names:
        if n.endswith('.xml') or n.endswith('.rels'):
            ET.fromstring(z.read(n))
    print(f'  XML well-formed: OK ({len([n for n in names if n.endswith((".xml", ".rels"))])} 个部件)')

    # 关系一致: workbook.xml 引用的 rId 必须存在于 workbook.xml.rels 且 Target 存在
    import posixpath
    rels_root = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    rid_map = {}
    for rel in rels_root:
        rid_map[rel.get('Id')] = rel.get('Target')
    used = set(re.findall(rb'r:id="(rId\d+)"', z.read('xl/workbook.xml')))
    for rid in used:
        rid = rid.decode()
        assert rid in rid_map, f'workbook.xml 引用了未定义的 {rid}'
        full = posixpath.normpath(posixpath.join('xl', rid_map[rid]))
        assert full in names, f'关系目标不存在: {rid_map[rid]}'
    print(f'  关系一致: OK ({len(used)} 处引用)')

    wb_src = json.loads(z.read('docProps/workbook.json'))
    assert 'sheets' in wb_src and len(wb_src['sheets']) > 0
    print(f'  内嵌 workbook.json: {len(wb_src["sheets"])} 个工作表')

    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == [s['name'] for s in wb_src['sheets']], f'表顺序不一致: {wb.sheetnames}'
    print(f'  openpyxl: 打开成功, 表顺序一致 {wb.sheetnames}')

    for si, s in enumerate(wb_src['sheets']):
        ws = wb.worksheets[si]
        headers = s.get('headers')
        rows = s.get('rows', [])
        n_rows = len(rows) + (1 if headers else 0)
        assert ws.max_row >= n_rows, f'{s["name"]}: 行数不足 {ws.max_row} < {n_rows}'
        # 表头加粗
        if headers:
            for c, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=c)
                assert cell.value == h, f'{s["name"]}: 表头 {c} 不一致 {cell.value!r} vs {h!r}'
                assert cell.font.bold, f'{s["name"]}: 表头未加粗: {h}'
            print(f'  {s["name"]}: 表头加粗 OK ({len(headers)} 列)')
        # 单元格值(数字与公式)
        for ri, row in enumerate(rows, start=1 + (1 if headers else 0)):
            for ci, v in enumerate(row, start=1):
                cell = ws.cell(row=ri, column=ci)
                if isinstance(v, str) and v.startswith('='):
                    got = cell.value
                    assert got == v, f'{s["name"]}!{cell.coordinate}: 公式 {got!r} != {v!r}'
                elif isinstance(v, (int, float)):
                    assert cell.value == v, f'{s["name"]}!{cell.coordinate}: 数字 {cell.value!r} != {v!r}'
        # 冻结窗格
        if s.get('freeze'):
            assert ws.freeze_panes == s['freeze'], f'{s["name"]}: 冻结窗格 {ws.freeze_panes} != {s["freeze"]}'
        # 数字格式(抽查第一列给了格式的)
        fmts = s.get('formats', [])
        if fmts and len(rows) > 0:
            f0 = fmts[0]
            if f0 in ('money', 'number'):
                assert '#,##0.00' in ws.cell(row=2, column=1).number_format, \
                    f'{s["name"]}: 第 1 列数字格式未生效 {ws.cell(row=2, column=1).number_format}'
            elif f0 == 'percent':
                assert '%' in ws.cell(row=2, column=1).number_format, \
                    f'{s["name"]}: 第 1 列百分比格式未生效'
            elif f0 == 'year':
                assert ws.cell(row=2, column=1).number_format in ('0', 'General'), \
                    f'{s["name"]}: 年份格式异常 {ws.cell(row=2, column=1).number_format}'
        # 列宽
        cw = s.get('colWidths')
        if cw and len(rows) > 0:
            assert abs(ws.column_dimensions['A'].width - cw[0]) < 0.5, \
                f'{s["name"]}: 列宽未生效 A={ws.column_dimensions["A"].width} vs {cw[0]}'
        print(f'  {s["name"]}: 值/冻结/格式/列宽核验 OK ({len(rows)} 数据行)')

def main():
    for path in sys.argv[1:]:
        validate(path)
    print('ALL VALIDATE OK')

if __name__ == '__main__':
    main()
