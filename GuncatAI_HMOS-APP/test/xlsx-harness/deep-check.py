# 深度校验: 内嵌 workbook.json 往返一致性 + openpyxl 读公式/数值/表顺序核验
# 用法: python deep-check.py gen/out_all.xlsx gen/out_md.xlsx gen/out_edit.xlsx
import sys
import zipfile
import json
from openpyxl import load_workbook

def deep(path):
    print(f'== {path} ==')
    z = zipfile.ZipFile(path)
    src = json.loads(z.read('docProps/workbook.json'))
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == [s['name'] for s in src['sheets']], '表顺序不一致'

    total_cells = 0
    checked_formulas = 0
    for si, s in enumerate(src['sheets']):
        ws = wb.worksheets[si]
        headers = s.get('headers', [])
        rows = s.get('rows', [])
        # 数据行逐格核验
        for ri, row in enumerate(rows, start=1 + (1 if headers else 0)):
            for ci, v in enumerate(row, start=1):
                cell = ws.cell(row=ri, column=ci)
                if isinstance(v, str) and v.startswith('='):
                    assert cell.value == v, f'{s["name"]}!{cell.coordinate} 公式不一致'
                    checked_formulas += 1
                elif isinstance(v, (int, float)):
                    assert cell.value == v, f'{s["name"]}!{cell.coordinate} 数值不一致'
                total_cells += 1
    # 公式数核验
    n_src_formula = sum(
        1 for s in src['sheets'] for row in s.get('rows', [])
        for v in row if isinstance(v, str) and v.startswith('=')
    )
    assert checked_formulas == n_src_formula, f'公式数不一致: {checked_formulas} vs {n_src_formula}'
    print(f'  内嵌源往返一致: {len(src["sheets"])} 表 / {total_cells} 格 / {n_src_formula} 个公式 OK')

    # 冻结窗格核验
    for s in src['sheets']:
        if s.get('freeze'):
            ws = wb[s['name']]
            assert ws.freeze_panes == s['freeze'], f'{s["name"]} 冻结窗格不一致'

if __name__ == '__main__':
    for path in sys.argv[1:]:
        deep(path)
    print('ALL DEEP-CHECK OK')
